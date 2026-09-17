import os
import openpyxl
from database import engine, init_db
from sqlmodel import Session
from models import Item

def parse_balance(val):
    if val is None:
        return None
    val_str = str(val).replace(',', '').strip()
    if val_str.startswith('(-)') or val_str.startswith('-'):
        num_str = val_str.replace('(-)', '').replace('-', '')
        try:
            return -abs(float(num_str))
        except:
            return None
    try:
        return float(val_str)
    except:
        return None

def run_import():
    db_path = "../data/inventory.db"
    if os.path.exists(db_path):
        print("Database already exists. Skipping first-run import.")
        return

    print("Initializing Database...")
    init_db()

    excel_path = "../data/StkSum_Department_Wise.xlsx"
    if not os.path.exists(excel_path):
        print(f"Excel file not found at {excel_path}. Please place it there to import data.")
        return

    print("Loading Excel file...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    with Session(engine) as session:
        count = 0
        for sheet_name in wb.sheetnames:
            if sheet_name == "Department Summary":
                continue
            
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if len(row) < 4:
                    continue
                sno, page, item_name, closing_balance = row[0:4]
                
                if not item_name:
                    continue
                
                cb = parse_balance(closing_balance)
                
                item = Item(
                    department=sheet_name,
                    page=page if type(page) == int else None,
                    item_name=str(item_name).strip(),
                    closing_balance=cb
                )
                session.add(item)
                count += 1
                
        session.commit()
        print(f"✅ Import complete! Imported {count} items.")

if __name__ == "__main__":
    run_import()
